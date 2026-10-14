from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest

from app.ingestion.chunking import chunk_document
from app.ingestion.parsers import ParsedDocument, ParseError, Section, parse_file


def _minimal_pdf(pages: list[str]) -> bytes:
    """Builds a tiny valid PDF with one text line per page (no PDF writer dependency)."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return out


def test_markdown_headings_become_paths(tmp_path: Path) -> None:
    path = tmp_path / "leave_policy.md"
    path.write_text(
        "Intro line.\n# Leave\n## Casual leave\nTwelve days.\n```\n# not a heading\n```\n"
        "## Sick leave\nTen days.\n# Holidays\nTen public holidays.\n",
        encoding="utf-8",
    )
    doc = parse_file(path)
    assert doc.title == "leave policy"
    paths = [s.heading_path for s in doc.sections]
    assert paths == [(), ("Leave", "Casual leave"), ("Leave", "Sick leave"), ("Holidays",)]
    assert "# not a heading" in doc.sections[1].text


def test_html_skips_scripts_and_uses_title(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>Travel Policy</title><style>h1{}</style></head><body>"
        "<script>alert('x')</script><h1>Travel</h1><p>Book economy &amp; save.</p>"
        "<h2>Per diem</h2><table><tr><td>City</td><td>40</td></tr></table></body></html>",
        encoding="utf-8",
    )
    doc = parse_file(path)
    assert doc.title == "Travel Policy"
    assert doc.sections[0] == Section(("Travel",), "Book economy & save.")
    assert doc.sections[1].heading_path == ("Travel", "Per diem")
    assert "alert" not in doc.text and "h1{}" not in doc.text


def test_docx_headings_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "handbook.docx"
    document = docx.Document()
    document.add_heading("Benefits", level=1)
    document.add_paragraph("Health cover for all staff.")
    document.add_heading("Dental", level=2)
    document.add_paragraph("Two checkups a year.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Plan", "Limit"
    table.cell(1, 0).text, table.cell(1, 1).text = "Basic", "500"
    document.save(str(path))
    doc = parse_file(path)
    assert [s.heading_path for s in doc.sections] == [
        ("Benefits",),
        ("Benefits", "Dental"),
        ("Table 1",),
    ]
    assert doc.sections[2].text == "Plan | Limit\nBasic | 500"


def test_xlsx_rows_keyed_by_header(tmp_path: Path) -> None:
    path = tmp_path / "holidays.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "2026"
    sheet.append(["Date", "Holiday"])
    sheet.append(["2026-01-26", "Founders Day"])
    sheet.append([None, None])
    workbook.create_sheet("Empty")
    workbook.save(str(path))
    doc = parse_file(path)
    assert doc.sections == [Section(("2026",), "Date: 2026-01-26; Holiday: Founders Day")]


def test_pdf_pages_tracked(tmp_path: Path) -> None:
    path = tmp_path / "guide.pdf"
    path.write_bytes(_minimal_pdf(["Page one text", "Page two text"]))
    doc = parse_file(path)
    assert [(s.page, s.text) for s in doc.sections] == [(1, "Page one text"), (2, "Page two text")]


@pytest.mark.parametrize(
    ("name", "content", "reason"),
    [
        ("broken.pdf", b"%PDF-1.4 not really", "could not parse .pdf file"),
        ("broken.docx", b"not a zip", "could not parse .docx file"),
        ("broken.xlsx", b"not a zip", "could not parse .xlsx file"),
        ("empty.md", b"   \n", "no extractable text"),
        ("image.png", b"\x89PNG", "unsupported file type"),
    ],
)
def test_bad_files_raise_parse_error(
    tmp_path: Path, name: str, content: bytes, reason: str
) -> None:
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(ParseError, match=reason):
        parse_file(path)


def test_chunks_respect_sections_size_and_overlap() -> None:
    sentence = "Employees accrue leave every month of service. "
    doc = ParsedDocument(
        title="Leave",
        media_type="text/markdown",
        sections=[
            Section(("A",), sentence * 60),
            Section(("B",), "Short section."),
        ],
    )
    chunks = chunk_document(doc, max_chars=400, overlap_chars=80)
    assert all(len(c.text) <= 400 for c in chunks)
    assert chunks[-1].heading_path == ("B",) and chunks[-1].text == "Short section."
    a_chunks = [c for c in chunks if c.heading_path == ("A",)]
    assert len(a_chunks) > 1
    assert a_chunks[1].text.startswith("Employees") or a_chunks[1].text[:20] in a_chunks[0].text
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert chunks[0].embedding_text("Leave").startswith("Leave > A\n\n")


def test_unbroken_text_is_hard_split() -> None:
    doc = ParsedDocument("t", "text/plain", [Section((), "x" * 1000)])
    chunks = chunk_document(doc, max_chars=300, overlap_chars=0)
    assert "".join(c.text for c in chunks) == "x" * 1000
    assert all(len(c.text) <= 300 for c in chunks)
