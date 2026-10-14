"""File parsers that turn a supported file into heading-aware sections.

Every parser returns ``ParsedDocument``; any library failure is raised as
``ParseError`` so callers never see pypdf/python-docx/openpyxl exceptions.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

MAX_FILE_BYTES = 25 * 1024 * 1024


class ParseError(Exception):
    def __init__(self, path: Path | str, reason: str) -> None:
        super().__init__(f"{Path(path).name}: {reason}")
        self.reason = reason


@dataclass(frozen=True)
class Section:
    heading_path: tuple[str, ...]
    text: str
    page: int | None = None


@dataclass
class ParsedDocument:
    title: str
    media_type: str
    sections: list[Section] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections)


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


class _SectionBuilder:
    """Collects body text under the current heading path."""

    def __init__(self) -> None:
        self.sections: list[Section] = []
        self._path: list[tuple[int, str]] = []
        self._buffer: list[str] = []

    def heading(self, level: int, title: str) -> None:
        title = _clean(title)
        if not title:
            return
        self.flush()
        self._path = [(lvl, t) for lvl, t in self._path if lvl < level]
        self._path.append((level, title))

    def text(self, text: str) -> None:
        self._buffer.append(text)

    def flush(self, page: int | None = None) -> None:
        body = _clean("\n".join(self._buffer))
        self._buffer = []
        if body:
            path = tuple(t for _, t in self._path)
            self.sections.append(Section(path, body, page))


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def parse_markdown(raw: str) -> list[Section]:
    builder = _SectionBuilder()
    in_fence = False
    for line in raw.splitlines():
        if _MD_FENCE.match(line):
            in_fence = not in_fence
        match = None if in_fence else _MD_HEADING.match(line)
        if match:
            builder.heading(len(match.group(1)), match.group(2))
        else:
            builder.text(line)
    builder.flush()
    return builder.sections


def parse_plain_text(raw: str) -> list[Section]:
    body = _clean(raw)
    return [Section((), body)] if body else []


class _HtmlSections(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template", "svg", "head"}
    _BLOCK = {"p", "div", "li", "tr", "br", "section", "article", "table", "ul", "ol", "pre"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.builder = _SectionBuilder()
        self.title = ""
        self._skip_depth = 0
        self._heading_level: int | None = None
        self._heading_text: list[str] = []
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
            self._in_title = False
        if tag == "title":
            self._in_title = True
        elif re.fullmatch(r"h[1-6]", tag):
            self._heading_level = int(tag[1])
            self._heading_text = []
        elif tag in self._BLOCK:
            self.builder.text("\n")
        elif tag in {"td", "th"}:
            self.builder.text(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        if tag == "title":
            self._in_title = False
        elif self._heading_level and tag == f"h{self._heading_level}":
            self.builder.heading(self._heading_level, " ".join(self._heading_text))
            self._heading_level = None
        elif tag in self._BLOCK:
            self.builder.text("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        elif self._skip_depth:
            return
        elif self._heading_level:
            self._heading_text.append(data)
        else:
            self.builder.text(data)


def parse_html(raw: str) -> tuple[str, list[Section]]:
    parser = _HtmlSections()
    parser.feed(raw)
    parser.close()
    parser.builder.flush()
    return _clean(parser.title), parser.builder.sections


def _parse_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ParseError(path, "encrypted PDF is not supported")
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        body = _clean(page.extract_text() or "")
        if body:
            sections.append(Section((), body, page=number))
    return sections


_DOCX_HEADING = re.compile(r"^Heading (\d)$")


def _parse_docx(path: Path) -> list[Section]:
    import docx

    document = docx.Document(str(path))
    builder = _SectionBuilder()
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        match = _DOCX_HEADING.match(style or "")
        if style == "Title":
            builder.heading(1, paragraph.text)
        elif match:
            builder.heading(int(match.group(1)), paragraph.text)
        else:
            builder.text(paragraph.text)
    builder.flush()
    for index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        body = _clean("\n".join(rows))
        if body:
            builder.sections.append(Section((f"Table {index}",), body))
    return builder.sections


def _parse_xlsx(path: Path) -> list[Section]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        sections = []
        for sheet in workbook.worksheets:
            rows = [
                [("" if v is None else str(v)).strip() for v in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            header, *body = rows
            if body and all(header):
                # Header-keyed lines keep each value next to its column name in a chunk.
                lines = [
                    "; ".join(f"{h}: {v}" for h, v in zip(header, r, strict=False) if v)
                    for r in body
                ]
            else:
                lines = [" | ".join(r) for r in rows]
            sections.append(Section((sheet.title,), _clean("\n".join(lines))))
        return sections
    finally:
        workbook.close()


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


MEDIA_TYPES = {
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

_BINARY_PARSERS: dict[str, Callable[[Path], list[Section]]] = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
}


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in MEDIA_TYPES


def parse_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix not in MEDIA_TYPES:
        raise ParseError(path, f"unsupported file type '{suffix}'")
    size = path.stat().st_size
    if size > MAX_FILE_BYTES:
        raise ParseError(path, f"file is {size} bytes; limit is {MAX_FILE_BYTES}")
    title = path.stem.replace("_", " ").replace("-", " ").strip()
    try:
        if suffix in _BINARY_PARSERS:
            sections = _BINARY_PARSERS[suffix](path)
        elif suffix in {".html", ".htm"}:
            html_title, sections = parse_html(_read_text(path))
            title = html_title or title
        elif suffix in {".md", ".markdown"}:
            sections = parse_markdown(_read_text(path))
        else:
            sections = parse_plain_text(_read_text(path))
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001 - library errors vary; never leak them raw
        raise ParseError(path, f"could not parse {suffix} file ({type(exc).__name__})") from exc
    if not sections:
        raise ParseError(path, "no extractable text")
    return ParsedDocument(title=title, media_type=MEDIA_TYPES[suffix], sections=sections)
